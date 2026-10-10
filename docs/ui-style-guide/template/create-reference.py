from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).parent
URL='https://app.notion.com/p/3e1ddaf9707a80c6a313ce6e7497828b'
pages=[
('Application identity & palette',[
('p','Duplicate this brief for each new application. Replace bracketed fields. Keep the family defaults unless an exception is documented. The full source audit and detailed component rules remain in the linked Notion reference.'),
('table', [['Project field','Your application'],['Application / platform','[Name] / [macOS, web, other]'],['Primary workflow','[What the user comes here to accomplish]'],['Main content / navigation','[Document, timeline, table, canvas] / [Navigation model]'],['Owner / version / date','[Owner] / [Version] / [Review date]']]),
('h','Family identity'),
('p','Quiet, compact, content-first workspaces. Opaque surfaces, restrained lavender emphasis, clear small icons, thin borders, soft rectangular corners, and stable navigation. Product names, icons, content, and specialized data colors may differ. Handwritten document lettering is not the default UI font.'),
('table',[['Semantic role','Light','Dark'],['Surface','#FFFFFF','#232329'],['Accent','#6965DB','#A8A5FF'],['Selection','#E3E2FE','#4F4D6F'],['Hover','#EFEDFF','#37334F'],['Structural border','#EBEBEB','#36363F'],['On-selection text*','#030064','#E0DFFF']]),
('p','*Recommended legibility improvement; the other five roles are observed native core colors. Use foreground/background pairs. Preserve adaptive system primary and secondary text where appropriate. Do not substitute the system-indigo AccentColor asset for the custom violet theme.'),
]),
('Spacing, type & component geometry',[
('table',[['Role','Family default / implementation direction'],['Spacing vocabulary','0, 2, 4, 6, 8, 10, 12, 16, 20, 24, 32, 40 pt or CSS px'],['Control gaps','8–10; panel insets 12–16; section gaps 16–24'],['Control heights','24 small; 32 regular; 36 large; 40 extra-large'],['Corner radii','6 controls; 8 cards/panels; 12 large dialogs'],['Structure','1-point borders; overlays alone receive restrained shadows'],['UI typography','Native system font; web system or Assistant, chosen consistently'],['Type hierarchy','13 medium buttons; 14 semibold headers; 11 status labels'],['Navigation row','8 horizontal / 6 vertical inner padding; 1 outer vertical'],['Header reference','52 high; keep content/actions aligned and predictable']]),
('h','Choose a shell suited to the product'),
('p','Use a navigator, main content, and optional inspector only when the workflow needs them. A menu-bar utility can share the same controls without using three panes. ExcalidrawDesktop references: sidebar outer min/ideal/max 230/250/340; main detail minimum 480; docked inspector 280/350/400. These are source observations, not universal window requirements.'),
('h','Shared components'),
('p','Surface, ActionButton, IconButton, NavigationRow, MenuRow, TextField, Section, ContentCard, Inspector, StatusLabel, EmptyState, and DialogActions. Centralize color, geometry, density, and motion preferences. Screens select semantic roles instead of repeating raw values.'),
('h','Product decisions'),
('p','[Chosen shell and minimum window size]\n[Typography and density by platform]\n[Component library location]\n[Documented exceptions and reasons]'),
]),
('Screens, interaction & complete states',[
('table',[['Screen / flow','Components / content','Required states'],['[Primary workspace]','[Navigator, content, inspector]','[Normal, narrow, long names]'],['[Browser / search]','[Rows, cards, filter controls]','[Empty, results, no matches]'],['[Settings / forms]','[Sections, inputs, actions]','[Focus, invalid, disabled]'],['[Import / export / save]','[Progress, dialog actions]','[Busy, success, failure, recovery]'],['[Secondary surfaces]','[Menus, history, previews]','[Unavailable, selected, loading]']]),
('h','Interaction contract'),
('p','Specify default, hover, pressed, selected, keyboard focus, disabled, loading, and error states for each relevant component. Selection and keyboard focus need independent indicators. Use explanatory error text and a recovery action; color alone is insufficient. Provide accessible labels for icon-only buttons.'),
('p','Observed desktop references: shared button pressed opacity 0.7 and disabled opacity 0.4; sidebar ease-in-out transition 0.16 seconds; multiselection outline 1.5 points. Apply these with contrast and reduced-motion checks rather than assuming opacity is always safe.'),
('h','Behavior that must remain dependable'),
('p','Preserve native window dragging, menus, keyboard shortcuts, file pickers, and focus semantics. Keep content geometry stable while loading. Support long text, resizing, and opening or closing inspectors. Avoid decorative gradients, large rounded cards, and material effects unless a specific surface has a documented reason.'),
('h','State decisions'),
('p','[Focus indicator and error roles]\n[Loading and recovery copy]\n[Keyboard routes and accessible names]\n[Reduced-motion behavior and performance expectations]'),
]),
('Adoption, acceptance & implementation prompt',[
('h','Implementation sequence'),
('p','1. Map current screens and preserve established workflows.\n2. Introduce shared color, type, spacing, and geometry tokens.\n3. Apply shared buttons and rows, then navigation and headers.\n4. Adapt forms, cards, inspectors, and secondary screens.\n5. Review real application states after each meaningful change.'),
('table',[['Acceptance area','Evidence to record'],['Appearance','[Light/dark screenshots; paired-color contrast checks]'],['Layout','[Narrow window, large text, long names, populated content]'],['Interaction','[Hover, selection, focus, pressed, disabled, errors]'],['Accessibility','[Keyboard-only flow, labels, targets, reduced motion]'],['Resilience','[Empty/loading/failure/recovery and stable geometry]'],['Platform / performance','[Native transitions, responsive navigation and previews]']]),
('h','Copyable implementation prompt'),
('p','Apply the Quiet Lavender family specification to [application]. Preserve its workflows, product name, and icon. Use this brief’s paired light/dark semantic colors, 6/8/12 corner hierarchy, 32-point regular controls, 8–10 control gaps, 12–16 panel insets, and 16–24 section gaps. Keep native system typography or a deliberate system/Assistant web choice. Implement complete interaction and availability states. Keep content dominant, ordinary surfaces flat, and native behavior intact. Read the full reference for component details and distinguish observed implementation from recommended rules. Validate the running application with realistic data; document necessary exceptions.'),
('h','Reference and sign-off'),
('p','Full 26-section guideline, interactive specimen, tokens, CSS, and source evidence:\n'+URL),
('p','[Reviewer] / [Date] / [Remaining gaps] / [Accepted exceptions]\nA specimen or build alone does not establish application acceptance.'),
])]

doc=Document(); sec=doc.sections[0]
for sty in doc.styles:
 for border in list(sty.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
sec.top_margin=sec.bottom_margin=Inches(.58)
sec.left_margin=sec.right_margin=Inches(.7)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
normal.font.color.rgb=RGBColor.from_string('232329')
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',27),('Heading 1',20),('Heading 2',12)]:
 s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('6965DB')
 s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(8)
for idx,(title,blocks) in enumerate(pages):
 if idx:doc.add_page_break()
 doc.add_paragraph('QUIET LAVENDER / APPLICATION BRIEF', 'Subtitle')
 doc.add_heading(title,0 if idx==0 else 1)
 for kind,value in blocks:
  if kind=='p':doc.add_paragraph(value)
  elif kind=='h':doc.add_heading(value,2)
  else:
   tab=doc.add_table(rows=0,cols=len(value[0]));tab.autofit=False
   for ri,row in enumerate(value):
    cells=tab.add_row().cells
    for c,txt in zip(cells,row):
     c.text=txt
     for p in c.paragraphs:
      p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
      for run in p.runs:run.font.size=Pt(9);run.bold=ri==0
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E3E2FE' if ri==0 else ('F7F7FA' if ri%2 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shade)
    if ri==0:
     flag=OxmlElement('w:tblHeader');tab.rows[-1]._tr.get_or_add_trPr().append(flag)
    flag=OxmlElement('w:cantSplit');tab.rows[-1]._tr.get_or_add_trPr().append(flag)
   doc.add_paragraph().paragraph_format.space_after=Pt(0)
doc.core_properties.title='Quiet Lavender Application Brief'
doc.core_properties.author=''
doc.save(OUT/'Quiet-Lavender-Application-Brief.docx')
md=['<callout icon="📋" color="purple_bg">\n\tDuplicate this page for each application and replace the bracketed fields. The family defaults remain shared; document product-specific exceptions.\n</callout>','<mention-page url="'+URL+'">Full style guideline</mention-page>','<table_of_contents/>']
for title,blocks in pages:
 md.append('## '+title)
 for kind,v in blocks:
  if kind=='p':md.append(v)
  elif kind=='h':md.append('### '+v)
  else:md.append('<table header-row="true">\n'+'\n'.join('<tr>'+''.join('<td>'+s+'</td>' for s in row)+'</tr>' for row in v)+'\n</table>')
(OUT/'notion-application-brief.md').write_text('\n\n'.join(md))
print('Created four-page application brief and Notion template content.')
